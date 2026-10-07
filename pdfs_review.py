# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (c) 2026 pdfs-review contributors
"""Review PDF pages in folders and nested ZIP, 7z and RAR archives.

Set ROOT_FOLDER below and run RUN.bat, or pass --root on the command line.
"""

from __future__ import annotations

import argparse
import codecs
import json
import os
import shutil
import struct
import sys
import tempfile
import time
import zipfile
import zlib
from collections import Counter
from datetime import datetime
from math import isfinite
from pathlib import Path

import fitz
import py7zr
import rarfile
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from py7zr.io import Py7zIO, WriterFactory

# Configuration: edit this path for a routine run. CLI options override it.
ROOT_FOLDER = "PASTE_FOLDER_PATH_HERE"
PROJECT_DIR = Path(__file__).resolve().parent
ARCHIVE_SUFFIXES = {".zip", ".7z", ".rar"}
MAX_ARCHIVE_DEPTH = 20
MAX_ARCHIVE_MEMBERS = 100_000
MAX_MEMBER_BYTES = 50 * 1024**3
MAX_EXTRACTED_BYTES = 100 * 1024**3
EXCEL_MAX_DATA_ROWS = 1_048_570
TITLE_ROW, COLUMN_NUMBER_ROW, HEADER_ROW, DATA_START_ROW = 2, 4, 6, 7
FIRST_TABLE_COLUMN = 2
NARROW_WIDTH, DATA_WIDTH = 5.81640625, 12.54296875
HEADER_HEIGHT, DEFAULT_ROW_HEIGHT = 48.65, 14.5
BASE_RESULT_COLUMNS = [
    "pdf_id",
    "object_type",
    "leaf_level",
    "full_path",
    "file_name",
    "file_size_bytes",
    "in_archive",
    "archive_depth",
    "pdf_page_count",
    "page_number",
    "class",
    "text_chars",
    "img_cover",
    "vector_count",
    "page_read_ok",
    "text_extract_ok",
    "has_extractable_text",
    "error_stage",
    "page_error",
]
PDF_RESULT_COLUMNS = [
    "pdf_id",
    "leaf_level",
    "full_path",
    "file_name",
    "file_size_bytes",
    "in_archive",
    "archive_depth",
    "pdf_page_count",
    "readiness",
    "readable",
    "text_ready",
    "pages_read",
    "pages_with_text",
    "pages_without_text",
    "page_read_errors",
    "text_extraction_errors",
    "classification_errors",
    "encrypted",
    "password_required",
    "parser_repaired",
    "parser_warning_count",
    "parser_warnings",
    "header_present",
    "eof_present",
    "file_error",
]


def page_text_metrics(page: fitz.Page) -> tuple[int, bool]:
    """Extract once: classification counts characters; readiness requires non-whitespace."""
    flags = fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES
    blocks = page.get_text("dict", flags=flags)["blocks"]
    text_chars, meaningful = 0, False
    for block in blocks:
        if block.get("type") == 0:
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    text = span.get("text", "")
                    text_chars += len(text)
                    meaningful = meaningful or bool(text.strip())
    return text_chars, meaningful


def _union_area(rects: list[tuple[float, float, float, float]]) -> float:
    if not rects:
        return 0.0
    xs = sorted({x for x0, _y0, x1, _y1 in rects for x in (x0, x1)})
    area = 0.0
    for index in range(len(xs) - 1):
        x0, x1 = xs[index], xs[index + 1]
        if x1 <= x0:
            continue
        ys = [(y0, y1) for rx0, y0, rx1, y1 in rects if not (rx1 <= x0 or rx0 >= x1)]
        if not ys:
            continue
        ys.sort()
        y0, y1 = ys[0]
        ysum = 0.0
        for lower, upper in ys[1:]:
            if upper <= lower:
                continue
            if lower <= y1:
                y1 = max(y1, upper)
            else:
                ysum += max(0.0, y1 - y0)
                y0, y1 = lower, upper
        ysum += max(0.0, y1 - y0)
        area += (x1 - x0) * ysum
    return area


def classify_page(
    page: fitz.Page,
    *,
    text_threshold: int,
    img_cover_threshold: float,
    min_img_frac: float,
    vector_threshold: int,
    text_chars: int | None = None,
) -> tuple[str, int, float, int]:
    if text_chars is None:
        text_chars, _ = page_text_metrics(page)
    page_rect = page.rect
    page_area = max(1.0, page_rect.width * page_rect.height)

    image_rects: list[tuple[float, float, float, float]] = []
    for image_info in page.get_image_info():
        bbox = (fitz.Rect(image_info.get("bbox", (0, 0, 0, 0))) * page.rotation_matrix) & page_rect
        if bbox.is_empty or bbox.is_infinite:
            continue
        area = bbox.width * bbox.height
        if area / page_area >= min_img_frac:
            image_rects.append((bbox.x0, bbox.y0, bbox.x1, bbox.y1))

    image_cover = min(1.0, _union_area(image_rects) / page_area)
    has_text = text_chars >= text_threshold
    vector_count = len(page.get_drawings())
    vector_flag = vector_count > vector_threshold

    if image_cover >= img_cover_threshold and not has_text:
        klass = "scan"
    elif image_cover >= img_cover_threshold and has_text:
        klass = "ocr"
    elif has_text:
        klass = "digital"
    elif vector_flag:
        klass = "vector"
    else:
        klass = "other"
    return klass, text_chars, round(image_cover, 4), vector_count


def styled_cell(
    ws,
    value: object,
    *,
    font: Font,
    border: Border | None = None,
    fill: PatternFill | None = None,
    alignment: Alignment | None = None,
    number_format: str | None = None,
    literal: bool = False,
) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    if literal and isinstance(value, str):
        if len(value.encode("utf-16-le")) // 2 > 32767:
            raise ValueError("Excel cell text exceeds 32767 characters")
        cell.data_type = "s"
    cell.font = font
    if border:
        cell.border = border
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if number_format:
        cell.number_format = number_format
    return cell


def build_xlsx(
    jsonl_path: Path,
    output_path: Path,
    row_count: int,
    max_level: int,
    pdfs_path: Path | None = None,
    pdf_count: int = 0,
) -> None:
    if max(row_count, pdf_count) > EXCEL_MAX_DATA_ROWS:
        raise RuntimeError(
            f"Worksheet data exceeds the Excel limit of {EXCEL_MAX_DATA_ROWS:,} rows "
            f"(pdf_pages={row_count:,}, pdf_files={pdf_count:,})"
        )

    workbook = Workbook(write_only=True)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.calcId = 191029
    workbook.views[0].windowWidth = 22260
    workbook.views[0].windowHeight = 12650
    workbook.views[0].tabRatio = 928

    add_flat_sheet(
        workbook,
        jsonl_path,
        row_count,
        max_level,
        BASE_RESULT_COLUMNS,
        "pdf_pages",
        "Table: PDF page review",
    )
    if pdfs_path is not None:
        add_flat_sheet(
            workbook,
            pdfs_path,
            pdf_count,
            max_level,
            PDF_RESULT_COLUMNS,
            "pdf_files",
            "Table: PDF file readiness",
        )
    os.makedirs(fs_path(output_path.parent), exist_ok=True)
    descriptor, temporary_output = tempfile.mkstemp(
        prefix="pdfs-", suffix=".xlsx", dir=fs_path(output_path.parent)
    )
    os.close(descriptor)
    try:
        workbook.save(temporary_output)
        os.replace(temporary_output, fs_path(output_path))
    finally:
        if os.path.exists(temporary_output):
            os.unlink(temporary_output)


def add_flat_sheet(workbook, jsonl_path, row_count, max_level, columns, sheet_name, title):
    ws = workbook.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_view.zoomScaleNormal = 85
    ws.sheet_properties.tabColor = "FF002060"
    ws.sheet_format.defaultRowHeight = DEFAULT_ROW_HEIGHT
    ws.row_dimensions[HEADER_ROW].height = HEADER_HEIGHT
    ws.freeze_panes = None
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = "9"  # A4
    ws.page_margins.left = 0.7
    ws.page_margins.right = 0.7
    ws.page_margins.top = 0.75
    ws.page_margins.bottom = 0.75
    ws.page_margins.header = 0.3
    ws.page_margins.footer = 0.3

    level_columns = [f"level_{index}" for index in range(1, max_level + 1)]
    result_columns = [*level_columns, *columns]
    headers = ["№пп", *result_columns, "last"]
    last_column = FIRST_TABLE_COLUMN + len(headers) - 1
    ws.column_dimensions["A"].width = NARROW_WIDTH
    for column in range(FIRST_TABLE_COLUMN, last_column + 1):
        letter = get_column_letter(column)
        is_boundary = column in {FIRST_TABLE_COLUMN, last_column}
        header = headers[column - FIRST_TABLE_COLUMN]
        wide_columns = {
            "full_path": 55,
            "file_name": 24,
            "page_error": 24,
            "file_error": 24,
            "parser_warnings": 24,
            "readiness": 27,
        }
        ws.column_dimensions[letter].width = (
            NARROW_WIDTH if is_boundary else wide_columns.get(header, DATA_WIDTH)
        )

    thin = Side(style="thin", color=Color(indexed=64))
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    normal_font = Font(name="Calibri", size=11, color=Color(theme=1))
    title_font = Font(name="Calibri", size=11, bold=True, color=Color(theme=1), charset=204)
    main_header_font = Font(name="Calibri", size=11, bold=True, color=Color(theme=1), charset=204)
    boundary_header_font = Font(
        name="Calibri", size=11, bold=True, color=Color(theme=0), charset=204
    )
    blue_fill = PatternFill("solid", fgColor="FF00B0F0")
    black_fill = PatternFill("solid", fgColor=Color(theme=1))
    centered = Alignment(horizontal="center", vertical="center")
    wrapped_centered = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.append([])
    ws.append(
        [
            None,
            styled_cell(
                ws,
                title,
                font=title_font,
            ),
        ]
    )
    ws.append([])

    number_row: list[object] = [None]
    for column in range(FIRST_TABLE_COLUMN, last_column + 1):
        number_row.append(
            styled_cell(
                ws,
                f"=COLUMN()-COLUMN($A${COLUMN_NUMBER_ROW})",
                font=normal_font,
                border=border,
                alignment=centered,
            )
        )
    ws.append(number_row)
    ws.append([])

    header_row: list[object] = [None]
    for index, header in enumerate(headers):
        is_boundary = index in {0, len(headers) - 1}
        header_row.append(
            styled_cell(
                ws,
                header,
                font=boundary_header_font if is_boundary else main_header_font,
                border=border,
                fill=black_fill if is_boundary else blue_fill,
                alignment=wrapped_centered,
            )
        )
    ws.append(header_row)

    centered_names = {
        "leaf_level",
        "in_archive",
        "archive_depth",
        "page_number",
        "class",
        "text_chars",
        "img_cover",
        "vector_count",
    }
    with jsonl_path.open("r", encoding="utf-8") as handle:
        for row_number, line in enumerate(handle, start=1):
            record = json.loads(line)
            levels = record.get("levels", [])
            excel_row: list[object] = [None]
            excel_row.append(
                styled_cell(
                    ws,
                    f"=ROW()-ROW($B${HEADER_ROW})",
                    font=normal_font,
                    border=border,
                    alignment=centered,
                )
            )
            for level_index in range(max_level):
                excel_row.append(
                    styled_cell(
                        ws,
                        levels[level_index] if level_index < len(levels) else None,
                        literal=True,
                        font=normal_font,
                        border=border,
                    )
                )
            for name in columns:
                alignment = centered if name in centered_names else None
                number_format = "0.0000" if name == "img_cover" else None
                excel_row.append(
                    styled_cell(
                        ws,
                        record.get(name),
                        literal=True,
                        font=normal_font,
                        border=border,
                        alignment=alignment,
                        number_format=number_format,
                    )
                )
            excel_row.append(
                styled_cell(
                    ws,
                    1,
                    font=normal_font,
                    border=border,
                    alignment=centered,
                )
            )
            ws.append(excel_row)

    last_row = max(HEADER_ROW, DATA_START_ROW + row_count - 1)
    ws.auto_filter.ref = (
        f"{get_column_letter(FIRST_TABLE_COLUMN)}{HEADER_ROW}:"
        f"{get_column_letter(last_column)}{last_row}"
    )
    ws.sheet_view.selection[0].activeCell = "B6"
    ws.sheet_view.selection[0].sqref = "B6"


# Filesystem and archive names
def fs_path(path: str | Path) -> str:
    value = os.path.abspath(os.fspath(path))
    if os.name == "nt" and not value.startswith("\\\\?\\"):
        return "\\\\?\\UNC\\" + value[2:] if value.startswith("\\\\") else "\\\\?\\" + value
    return value


def suffix(name: str) -> str:
    return Path(name.replace("\\", "/")).suffix.casefold()


def selected(name: str) -> bool:
    return suffix(name) == ".pdf" or suffix(name) in ARCHIVE_SUFFIXES


def zip_name(info: zipfile.ZipInfo, encoding: str) -> str:
    if info.flag_bits & 0x800:
        return info.orig_filename
    raw = info.orig_filename.encode("cp437")
    offset = 0
    while offset + 4 <= len(info.extra):
        kind, size = struct.unpack_from("<HH", info.extra, offset)
        body = info.extra[offset + 4 : offset + 4 + size]
        if kind == 0x7075 and len(body) >= 5 and body[0] == 1:
            if struct.unpack_from("<I", body, 1)[0] == zlib.crc32(raw):
                return body[5:].decode("utf-8")
        offset += 4 + size
    return raw.decode(encoding)


def configure_rar_backend(explicit: str | None) -> None:
    candidates = [explicit, shutil.which("7z"), shutil.which("7zz")]
    if os.name == "nt":
        candidates.append(
            str(Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "7-Zip" / "7z.exe")
        )
    for candidate in candidates:
        if candidate and os.path.isfile(candidate):
            rarfile.SEVENZIP_TOOL = candidate
            return


class Review:
    def __init__(self, args, rows, errors, temp_root: Path, pdfs=None):
        self.args, self.rows, self.errors, self.temp_root = args, rows, errors, temp_root
        self.pdfs = pdfs
        self.stats = Counter()
        self.classes = Counter()
        self.readiness = Counter()
        self.last_progress = time.monotonic()

    def progress(self):
        if time.monotonic() - self.last_progress >= 10:
            print(
                f"PDFs: {self.stats['pdfs_found']:,}; pages: {self.stats['pages']:,}; "
                f"archives: {self.stats['archives']:,}; skipped: {self.stats['skipped_files']:,}; "
                f"errors: {self.stats['errors']:,}",
                flush=True,
            )
            self.last_progress = time.monotonic()

    def base(self, logical, levels, depth, size=None, object_type="pdf"):
        return {
            "levels": list(levels),
            "leaf_level": len(levels),
            "full_path": logical,
            "file_name": levels[-1] if levels else None,
            "file_size_bytes": size,
            "in_archive": "yes" if depth else "no",
            "archive_depth": depth,
            "object_type": object_type,
            "pdf_id": None,
            "pdf_page_count": None,
        }

    def emit(
        self,
        base,
        page=None,
        klass="error",
        text=None,
        cover=None,
        vectors=None,
        error="",
        **checks,
    ):
        self.rows.write(
            json.dumps(
                {
                    **base,
                    "page_number": page,
                    "class": klass,
                    "text_chars": text,
                    "img_cover": cover,
                    "vector_count": vectors,
                    "page_error": error,
                    **checks,
                },
                ensure_ascii=False,
            )
            + "\n"
        )
        self.stats["rows"] += 1
        self.stats["max_level"] = max(self.stats["max_level"], len(base["levels"]))
        if page is not None:
            self.stats["pages"] += 1
            self.classes[klass] += 1
        self.progress()

    def failure(self, base, exc, page=None, **checks):
        new_pdf = base["object_type"] == "pdf" and base["pdf_id"] is None
        if base["object_type"] == "pdf" and base["pdf_id"] is None:
            self.stats["pdfs_found"] += 1
            base["pdf_id"] = self.stats["pdfs_found"]
        message = " ".join(f"{type(exc).__name__}: {exc}".splitlines())
        self.stats["errors"] += 1
        self.errors.write(f"{base['full_path']}\t{page or ''}\t{message}\n")
        self.errors.flush()
        self.emit(base, page=page, error=message, **checks)
        if new_pdf:
            self.finish_pdf({**base, "file_error": message})
        return message

    def finish_pdf(self, record):
        total = record.get("pdf_page_count")
        readable = bool(
            total and record.get("pages_read") == total and not record.get("file_error")
        )
        text_ok = readable and not record.get("text_extraction_errors")
        if not readable:
            readiness = "PDF_NOT_READY"
        elif not text_ok:
            readiness = "READY_TEXT_PARTIAL"
        elif record.get("pages_with_text", 0):
            readiness = "READY_WITH_TEXT"
        else:
            readiness = "READY_NO_TEXT_LAYER"
        record.update(
            readiness=readiness,
            readable=readable,
            text_ready=bool(text_ok and record.get("pages_with_text", 0)),
        )
        if self.pdfs is not None:
            self.pdfs.write(json.dumps(record, ensure_ascii=False) + "\n")
        self.readiness[readiness] += 1
        self.stats["pdf_file_rows"] += 1
        self.stats["pdfs_readable"] += int(readable)
        self.stats["pdfs_text_ready"] += int(record["text_ready"])
        self.stats["parser_warnings"] += record.get("parser_warning_count", 0)

    def check_size(self, size):
        if size > MAX_MEMBER_BYTES:
            raise ValueError("Archive member exceeds the configured size limit")
        if self.stats["decompressed_bytes"] + size > self.args.max_extracted_gib * 1024**3:
            raise ValueError("Cumulative archive decompression limit reached")

    def consume(self, size):
        self.check_size(size)
        self.stats["decompressed_bytes"] += size

    def copy_member(self, source, target, expected):
        self.check_size(expected)
        written = 0
        with open(fs_path(target), "wb") as output:
            while chunk := source.read(1024 * 1024):
                written += len(chunk)
                if written > MAX_MEMBER_BYTES:
                    raise ValueError("Archive member exceeds the configured size limit")
                self.consume(len(chunk))
                output.write(chunk)
        if written != expected:
            raise ValueError("Archive member size mismatch")

    def pdf(self, path, logical, levels, depth):
        self.stats["pdfs_found"] += 1
        base = self.base(logical, levels, depth)
        base["pdf_id"] = self.stats["pdfs_found"]
        record = {
            **base,
            "pages_read": 0,
            "pages_with_text": 0,
            "pages_without_text": 0,
            "page_read_errors": 0,
            "text_extraction_errors": 0,
            "classification_errors": 0,
            "file_error": "",
        }
        old_errors = fitz.TOOLS.mupdf_display_errors()
        old_warnings = fitz.TOOLS.mupdf_display_warnings()
        fitz.TOOLS.mupdf_display_errors(False)
        fitz.TOOLS.mupdf_display_warnings(False)
        fitz.TOOLS.reset_mupdf_warnings()
        try:
            base["file_size_bytes"] = os.stat(fs_path(path)).st_size
            if base["file_size_bytes"] > self.args.max_pdf_gib * 1024**3:
                raise ValueError("PDF exceeds the configured size limit")
            with open(fs_path(path), "rb") as source:
                record["header_present"] = b"%PDF-" in source.read(1024)
                source.seek(max(0, base["file_size_bytes"] - 4096))
                record["eof_present"] = b"%%EOF" in source.read()
            with fitz.open(fs_path(path), filetype="pdf") as doc:
                if not doc.is_pdf:
                    raise ValueError("File is not a PDF")
                record["encrypted"] = bool(
                    doc.is_encrypted or (doc.metadata or {}).get("encryption")
                )
                record["password_required"] = bool(doc.needs_pass and not doc.authenticate(""))
                record["parser_repaired"] = bool(doc.is_repaired)
                if record["password_required"]:
                    raise ValueError("PDF requires a password")
                base["pdf_page_count"] = doc.page_count
                if not doc.page_count:
                    raise ValueError("PDF has no pages")
                if doc.page_count > self.args.max_pages:
                    raise ValueError("PDF exceeds the configured page limit")
                self.stats["pdfs_opened"] += 1
                for index in range(doc.page_count):
                    try:
                        page = doc[index]
                        if page.rect.is_empty or page.rect.is_infinite:
                            raise ValueError("Page has invalid dimensions")
                    except Exception as exc:
                        record["page_read_errors"] += 1
                        self.stats["page_read_errors"] += 1
                        self.failure(
                            base, exc, page=index + 1, error_stage="page_read", page_read_ok=False
                        )
                        continue
                    record["pages_read"] += 1
                    try:
                        text_chars, meaningful = page_text_metrics(page)
                    except Exception as exc:
                        record["text_extraction_errors"] += 1
                        self.stats["text_extraction_errors"] += 1
                        self.failure(
                            base,
                            exc,
                            page=index + 1,
                            error_stage="text_extraction",
                            page_read_ok=True,
                            text_extract_ok=False,
                        )
                        continue
                    key = "pages_with_text" if meaningful else "pages_without_text"
                    record[key] += 1
                    self.stats[key] += 1
                    checks = dict(
                        page_read_ok=True, text_extract_ok=True, has_extractable_text=meaningful
                    )
                    try:
                        result = classify_page(
                            page,
                            text_chars=text_chars,
                            text_threshold=self.args.text_threshold,
                            img_cover_threshold=self.args.img_cover_threshold,
                            min_img_frac=self.args.min_img_frac,
                            vector_threshold=self.args.vector_threshold,
                        )
                    except Exception as exc:
                        record["classification_errors"] += 1
                        self.stats["classification_errors"] += 1
                        self.failure(
                            base,
                            exc,
                            page=index + 1,
                            error_stage="classification",
                            text=text_chars,
                            **checks,
                        )
                    else:
                        self.emit(base, index + 1, *result, **checks)
        except Exception as exc:
            record["file_error"] = self.failure(base, exc, error_stage="pdf_open_or_limit")
        finally:
            diagnostics = fitz.TOOLS.mupdf_warnings()
            fitz.TOOLS.mupdf_display_errors(old_errors)
            fitz.TOOLS.mupdf_display_warnings(old_warnings)
            record.update(base)
            record["parser_warnings"] = diagnostics
            record["parser_warning_count"] = len(diagnostics.splitlines())
            self.finish_pdf(record)

    def member(self, source, info_size, name, logical, levels, depth, temporary):
        parts = tuple(part for part in name.replace("\\", "/").split("/") if part)
        child_levels = levels + parts
        child_logical = logical + "::" + name
        base = self.base(
            child_logical,
            child_levels,
            depth,
            info_size,
            "archive" if suffix(name) in ARCHIVE_SUFFIXES else "pdf",
        )
        path = Path(temporary) / "member.bin"
        copied = False
        try:
            self.copy_member(source, path, info_size)
        except Exception as exc:
            self.failure(base, exc)
        else:
            copied = True
            self.visit(path, child_logical, child_levels, depth)
        finally:
            if os.path.exists(fs_path(path)):
                os.unlink(fs_path(path))
        return copied

    def drain_member(self, source, expected):
        """Consume to EOF for ZIP CRC validation without retaining unrelated files."""
        self.check_size(expected)
        written = 0
        while chunk := source.read(1024 * 1024):
            written += len(chunk)
            if written > MAX_MEMBER_BYTES:
                raise ValueError("Archive member exceeds the configured size limit")
            self.consume(len(chunk))
        if written != expected:
            raise ValueError("Archive member size mismatch")

    def zip_or_rar(self, path, logical, levels, depth, temporary, kind):
        archive_class = zipfile.ZipFile if kind == ".zip" else rarfile.RarFile
        with archive_class(fs_path(path)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_MEMBERS:
                raise ValueError("Archive member count exceeds the configured limit")
            full_crc = kind == ".zip" and self.args.check_zip_crc
            crc_ok = True
            if kind == ".zip":
                self.stats["zips_opened"] += 1
                has_pdf = any(
                    not i.is_dir() and suffix(zip_name(i, self.args.zip_encoding)) == ".pdf"
                    for i in infos
                )
                self.stats["zips_with_direct_pdf" if has_pdf else "zips_without_direct_pdf"] += 1
            for info in infos:
                if info.is_dir():
                    continue
                if kind == ".rar" and info.is_symlink():
                    self.stats["skipped_files"] += 1
                    continue
                name = zip_name(info, self.args.zip_encoding) if kind == ".zip" else info.filename
                if not selected(name):
                    self.stats["skipped_files"] += 1
                    if not full_crc:
                        continue
                try:
                    source = archive.open(info)
                except Exception as exc:
                    crc_ok = False
                    parts = tuple(p for p in name.replace("\\", "/").split("/") if p)
                    self.failure(
                        self.base(
                            logical + "::" + name,
                            levels + parts,
                            depth,
                            info.file_size,
                            "archive"
                            if suffix(name) in ARCHIVE_SUFFIXES
                            else "pdf"
                            if suffix(name) == ".pdf"
                            else "file",
                        ),
                        exc,
                    )
                else:
                    with source:
                        if selected(name):
                            checked = self.member(
                                source, info.file_size, name, logical, levels, depth, temporary
                            )
                        else:
                            try:
                                self.drain_member(source, info.file_size)
                            except Exception as exc:
                                parts = tuple(p for p in name.replace("\\", "/").split("/") if p)
                                self.failure(
                                    self.base(
                                        logical + "::" + name,
                                        levels + parts,
                                        depth,
                                        info.file_size,
                                        "file",
                                    ),
                                    exc,
                                )
                                checked = False
                            else:
                                checked = True
                        crc_ok = crc_ok and checked
                        if kind == ".zip" and checked:
                            self.stats["zip_members_crc_checked"] += 1
            if full_crc:
                self.stats["zips_full_crc_passed" if crc_ok else "zips_full_crc_failed"] += 1

    def seven_zip(self, path, logical, levels, depth, temporary):
        with py7zr.SevenZipFile(fs_path(path), mode="r") as archive:
            if archive.needs_password():
                raise ValueError("7z archive requires a password")
            infos = [info for info in archive.list() if not info.is_directory]
            if len(infos) > MAX_ARCHIVE_MEMBERS:
                raise ValueError("Archive member count exceeds the configured limit")
            selected_infos = [
                info for info in infos if selected(info.filename) and not info.is_symlink
            ]
            self.stats["skipped_files"] += len(infos) - len(selected_infos)
            if not selected_infos:
                return
            names = [info.filename for info in selected_infos]
            if len(names) != len(set(names)):
                raise ValueError("Duplicate selected 7z member names are unsupported")
            for info in infos:
                self.check_size(info.uncompressed or 0)
            total = sum(info.uncompressed or 0 for info in infos)
            if self.stats["decompressed_bytes"] + total > self.args.max_extracted_gib * 1024**3:
                raise ValueError("Cumulative archive decompression limit reached")
            factory = CaptureFactory(self, {i.filename for i in selected_infos}, temporary)
            try:
                # One decoder pass also for solid archives. Unrelated members are discarded.
                archive.extractall(factory=factory)
            finally:
                factory.close()
            for info in selected_infos:
                product = factory.products.get(info.filename)
                if product is None or product.size() != (info.uncompressed or 0):
                    raise ValueError("7z member was not captured completely")
                parts = tuple(p for p in info.filename.replace("\\", "/").split("/") if p)
                try:
                    self.visit(product.path, logical + "::" + info.filename, levels + parts, depth)
                finally:
                    os.unlink(fs_path(product.path))

    def visit(self, path, logical, levels, depth):
        kind = suffix(levels[-1])
        if kind == ".pdf":
            self.pdf(path, logical, levels, depth)
        elif kind in ARCHIVE_SUFFIXES:
            self.stats["archives"] += 1
            if kind == ".zip":
                self.stats["zips_found"] += 1
            try:
                if depth >= self.args.max_depth:
                    raise ValueError("Archive nesting limit reached")
                with tempfile.TemporaryDirectory(
                    prefix="archive-", dir=self.temp_root
                ) as temporary:
                    if kind == ".7z":
                        self.seven_zip(path, logical, levels, depth + 1, temporary)
                    else:
                        self.zip_or_rar(path, logical, levels, depth + 1, temporary, kind)
            except Exception as exc:
                self.failure(self.base(logical, levels, depth, object_type="archive"), exc)
        else:
            self.stats["skipped_files"] += 1
        self.progress()

    def scan(self, root):
        pending = [(root, ())]
        while pending:
            directory, prefix = pending.pop()
            try:
                with os.scandir(fs_path(directory)) as iterator:
                    entries = sorted(iterator, key=lambda item: (item.name.casefold(), item.name))
            except OSError as exc:
                self.failure(self.base(str(directory), prefix, 0, object_type="folder"), exc)
                continue
            children = []
            for entry in entries:
                path = directory / entry.name
                if path == self.args.output or path == self.args.runtime:
                    continue
                levels = prefix + (entry.name,)
                try:
                    if entry.is_symlink() or (
                        hasattr(entry, "is_junction") and entry.is_junction()
                    ):
                        continue
                    if entry.is_dir(follow_symlinks=False):
                        children.append((path, levels))
                    elif entry.is_file(follow_symlinks=False):
                        self.visit(path, str(path), levels, 0)
                except OSError as exc:
                    self.failure(self.base(str(path), levels, 0, object_type="file"), exc)
            pending.extend(reversed(children))


class Capture(Py7zIO):
    def __init__(self, review, path=None):
        self.review, self.path, self.written = review, path, 0
        self.handle = open(fs_path(path), "w+b") if path else None

    def write(self, chunk):
        self.written += len(chunk)
        if self.written > MAX_MEMBER_BYTES:
            raise ValueError("Archive member exceeds the configured size limit")
        self.review.consume(len(chunk))
        if self.handle:
            self.handle.write(chunk)
        return len(chunk)

    def read(self, size=None):
        return self.handle.read(-1 if size is None else size) if self.handle else b""

    def seek(self, offset, whence=0):
        return self.handle.seek(offset, whence) if self.handle else 0

    def flush(self):
        if self.handle:
            self.handle.flush()

    def size(self):
        return self.written

    def close(self):
        if self.handle:
            self.handle.close()
            self.handle = None


class CaptureFactory(WriterFactory):
    def __init__(self, review, selected_names, temporary):
        self.review, self.selected_names, self.temporary = review, selected_names, temporary
        self.products = {}

    def create(self, filename):
        path = (
            Path(self.temporary) / f"{len(self.products)}.bin"
            if filename in self.selected_names
            else None
        )
        product = Capture(self.review, path)
        self.products[filename] = product
        return product

    def close(self):
        for product in self.products.values():
            product.close()


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description="Review PDF pages in folders and nested ZIP/7z/RAR archives."
    )
    parser.add_argument("--root", default=ROOT_FOLDER, help="Input folder (overrides ROOT_FOLDER)")
    parser.add_argument("--output", type=Path, default=PROJECT_DIR / "pdfs-review.xlsx")
    parser.add_argument(
        "--zip-encoding", default="cp437", help="Unmarked ZIP names only; e.g. cp866 or cp1251"
    )
    parser.add_argument("--seven-zip", help="Optional 7z executable path for RAR support")
    parser.add_argument("--max-depth", type=int, default=MAX_ARCHIVE_DEPTH)
    parser.add_argument("--max-extracted-gib", type=float, default=MAX_EXTRACTED_BYTES / 1024**3)
    parser.add_argument("--max-pdf-gib", type=float, default=2)
    parser.add_argument("--max-pages", type=int, default=10_000)
    parser.add_argument(
        "--check-zip-crc",
        action="store_true",
        help="Read and CRC-check every ZIP member, including skipped extensions",
    )
    parser.add_argument(
        "--timestamp-output",
        action="store_true",
        help="Add a timestamp to the output name to retain previous workbooks",
    )
    parser.add_argument("--text-threshold", type=int, default=40)
    parser.add_argument("--img-cover-threshold", type=float, default=0.90)
    parser.add_argument("--min-img-frac", type=float, default=0.005)
    parser.add_argument("--vector-threshold", type=int, default=100)
    return parser.parse_args(argv)


def run(argv=None):
    args = parse_args(argv)
    if args.root == "PASTE_FOLDER_PATH_HERE":
        raise ValueError("Set ROOT_FOLDER in pdfs_review.py or pass --root")
    root = Path(os.path.abspath(args.root))
    args.output = Path(os.path.abspath(args.output))
    if args.timestamp_output:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        args.output = args.output.with_name(f"{args.output.stem}-{stamp}{args.output.suffix}")
    args.runtime = PROJECT_DIR / ".runtime"
    if not os.path.isdir(fs_path(root)):
        raise ValueError("Input folder is unavailable")
    if args.output.suffix.casefold() != ".xlsx":
        raise ValueError("Output must end with .xlsx")
    limits = (args.max_depth, args.max_extracted_gib, args.max_pdf_gib, args.max_pages)
    if any(not isfinite(value) or value <= 0 for value in limits):
        raise ValueError("Archive and PDF limits must be positive")
    if args.text_threshold < 0 or args.vector_threshold < 0:
        raise ValueError("Text and vector thresholds must be nonnegative")
    if not (0 <= args.img_cover_threshold <= 1 and 0 <= args.min_img_frac <= 1):
        raise ValueError("Image thresholds must be between 0 and 1")
    codecs.lookup(args.zip_encoding)
    configure_rar_backend(args.seven_zip)
    args.runtime.mkdir(exist_ok=True)
    spool = args.runtime / "pages.jsonl"
    pdf_spool = args.runtime / "pdfs.jsonl"
    errors = args.output.with_name(args.output.stem + "-errors.txt")
    summary = args.output.with_name(args.output.stem + "-summary.json")
    os.makedirs(fs_path(args.output.parent), exist_ok=True)
    started = time.monotonic()
    print(f"Input: {root}\nOutput: {args.output}", flush=True)
    with tempfile.TemporaryDirectory(prefix="review-", dir=args.runtime) as temporary:
        with (
            spool.open("w", encoding="utf-8") as rows,
            pdf_spool.open("w", encoding="utf-8") as pdfs,
            open(fs_path(errors), "w", encoding="utf-8") as log,
        ):
            review = Review(args, rows, log, Path(temporary), pdfs)
            review.scan(root)
        if review.stats["pdf_file_rows"] != review.stats["pdfs_found"]:
            raise RuntimeError("PDF discovery/file overview count mismatch")
        print("Writing workbook...", flush=True)
        build_xlsx(
            spool,
            args.output,
            review.stats["rows"],
            review.stats["max_level"],
            pdf_spool,
            review.stats["pdf_file_rows"],
        )
    workbook = load_workbook(fs_path(args.output), read_only=True)
    try:
        for name, key in (("pdf_pages", "rows"), ("pdf_files", "pdf_file_rows")):
            written_rows = sum(1 for _ in workbook[name].iter_rows(min_row=7, values_only=True))
            if written_rows != review.stats[key]:
                raise RuntimeError(f"Workbook {name} row count mismatch")
    finally:
        workbook.close()
    keys = (
        "pdfs_found",
        "pdfs_opened",
        "pages",
        "rows",
        "archives",
        "skipped_files",
        "decompressed_bytes",
        "errors",
        "max_level",
        "pdf_file_rows",
        "pdfs_readable",
        "pdfs_text_ready",
        "pages_with_text",
        "pages_without_text",
        "page_read_errors",
        "text_extraction_errors",
        "classification_errors",
        "parser_warnings",
        "zips_opened",
        "zips_with_direct_pdf",
        "zips_found",
        "zips_without_direct_pdf",
        "zip_members_crc_checked",
        "zips_full_crc_passed",
        "zips_full_crc_failed",
    )
    result = {key: review.stats[key] for key in keys}
    result.update(
        {
            "classes": dict(review.classes),
            "readiness": dict(review.readiness),
            "zip_full_crc_requested": args.check_zip_crc,
            "elapsed_seconds": round(time.monotonic() - started, 1),
        }
    )
    with open(fs_path(summary), "w", encoding="utf-8") as handle:
        json.dump(result, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    print(
        f"Done: {result['pdfs_found']:,} PDFs; {result['pages']:,} pages; "
        f"{result['skipped_files']:,} skipped files; {result['errors']:,} errors.\n"
        f"Workbook: {args.output}\nSummary: {summary}\nErrors: {errors}",
        flush=True,
    )
    return 2 if review.stats["errors"] else 0


def main():
    try:
        return run()
    except KeyboardInterrupt:
        print("Interrupted. Rerun to start a fresh review.", file=sys.stderr)
        return 130
    except Exception as exc:
        print(f"ERROR: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
