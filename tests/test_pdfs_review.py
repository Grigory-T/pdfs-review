from __future__ import annotations

import io
import json
import os
import struct
import tempfile
import unittest
import warnings
import zipfile
import zlib
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

import fitz
import py7zr
from openpyxl import load_workbook

import pdfs_review as app


def sample_pdf(pages=1):
    with fitz.open() as doc:
        for _ in range(pages):
            page = doc.new_page()
            page.insert_text(
                (30, 40), "This page contains more than forty extracted text characters."
            )
        return doc.tobytes()


def zip_bytes(members):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members:
            archive.writestr(name, payload)
    return stream.getvalue()


def rar_bytes(members):
    """Create a neutral RAR4 STORE archive; no external RAR creation utility."""

    def header(body):
        return struct.pack("<H", zlib.crc32(body) & 0xFFFF) + body

    result = b"Rar!\x1a\x07\x00" + header(struct.pack("<BHHHI", 0x73, 0, 13, 0, 0))
    for name, data in members:
        raw = name.encode("ascii")
        body = (
            struct.pack(
                "<BHHIIBIIBBHI",
                0x74,
                0x8000,
                32 + len(raw),
                len(data),
                len(data),
                3,
                zlib.crc32(data),
                0,
                20,
                0x30,
                len(raw),
                0o100644,
            )
            + raw
        )
        result += header(body) + data
    return result + header(struct.pack("<BHH", 0x7B, 0, 7))


def legacy_zip(name, data, encoding, unicode_name=None):
    raw = name.encode(encoding)
    extra = b""
    if unicode_name:
        field = b"\x01" + struct.pack("<I", zlib.crc32(raw)) + unicode_name.encode("utf-8")
        extra = struct.pack("<HH", 0x7075, len(field)) + field
    crc = zlib.crc32(data)
    local = (
        struct.pack(
            "<IHHHHHIIIHH",
            0x04034B50,
            20,
            0,
            0,
            0,
            0,
            crc,
            len(data),
            len(data),
            len(raw),
            len(extra),
        )
        + raw
        + extra
        + data
    )
    central = (
        struct.pack(
            "<IHHHHHHIIIHHHHHII",
            0x02014B50,
            20,
            20,
            0,
            0,
            0,
            0,
            crc,
            len(data),
            len(data),
            len(raw),
            len(extra),
            0,
            0,
            0,
            0,
            0,
        )
        + raw
        + extra
    )
    return (
        local
        + central
        + struct.pack("<IHHHHIIH", 0x06054B50, 0, 0, 1, 1, len(central), len(local), 0)
    )


class ReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="pdfs-review-test-")
        self.root = Path(self.temp.name)
        self.inputs = self.root / "inputs"
        self.inputs.mkdir()

    def tearDown(self):
        self.temp.cleanup()

    def seven_zip(self, members):
        stream = io.BytesIO()
        with py7zr.SevenZipFile(stream, "w") as archive:
            for name, payload in members:
                archive.writestr(payload, name)
        return stream.getvalue()

    def scan(self, **overrides):
        args = app.parse_args(["--root", str(self.inputs)])
        for key, value in overrides.items():
            setattr(args, key, value)
        args.output = self.root / "result.xlsx"
        args.runtime = self.root / "runtime"
        args.runtime.mkdir(exist_ok=True)
        rows, errors = io.StringIO(), io.StringIO()
        app.configure_rar_backend(None)
        review = app.Review(args, rows, errors, args.runtime, io.StringIO())
        review.scan(self.inputs)
        records = [json.loads(line) for line in rows.getvalue().splitlines()]
        self.assertEqual(list(args.runtime.iterdir()), [])
        return records, review

    def files(self, review):
        return [json.loads(line) for line in review.pdfs.getvalue().splitlines()]

    def test_readiness_is_independent_of_classification_threshold(self):
        with fitz.open() as doc:
            doc.new_page().insert_text((30, 40), "Hi")
            doc.new_page()
            (self.inputs / "short.pdf").write_bytes(doc.tobytes())
        with fitz.open() as doc:
            doc.new_page()
            (self.inputs / "blank.pdf").write_bytes(doc.tobytes())
        rows, review = self.scan()
        files = {r["file_name"]: r for r in self.files(review)}
        self.assertEqual(files["short.pdf"]["readiness"], "READY_WITH_TEXT")
        self.assertEqual(files["short.pdf"]["pages_with_text"], 1)
        self.assertEqual(files["short.pdf"]["pages_without_text"], 1)
        self.assertEqual(files["blank.pdf"]["readiness"], "READY_NO_TEXT_LAYER")
        self.assertTrue(all(r["class"] == "other" for r in rows))
        self.assertEqual(sum(r["has_extractable_text"] for r in rows), 1)
        self.assertEqual(review.stats["errors"], 0)

    def test_text_failure_is_partial_not_unreadable(self):
        (self.inputs / "two.pdf").write_bytes(sample_pdf(2))
        original = app.page_text_metrics

        def extract(page):
            if page.number == 0:
                raise ValueError("synthetic text failure")
            return original(page)

        with mock.patch.object(app, "page_text_metrics", side_effect=extract):
            rows, review = self.scan()
        file = self.files(review)[0]
        self.assertEqual(file["readiness"], "READY_TEXT_PARTIAL")
        self.assertTrue(file["readable"])
        self.assertFalse(file["text_ready"])
        self.assertEqual(file["text_extraction_errors"], 1)
        self.assertEqual(rows[0]["error_stage"], "text_extraction")
        self.assertTrue(rows[0]["page_read_ok"])
        self.assertFalse(rows[0]["text_extract_ok"])
        self.assertIsNone(rows[0].get("has_extractable_text"))
        self.assertEqual(rows[1]["class"], "digital")

    def test_page_read_failure_does_not_mark_whole_pdf_readable(self):
        (self.inputs / "two.pdf").write_bytes(sample_pdf(2))
        original = fitz.Document.load_page

        def load(doc, index):
            if index == 0:
                raise ValueError("synthetic page read failure")
            return original(doc, index)

        with mock.patch.object(fitz.Document, "load_page", load):
            rows, review = self.scan()
        file = self.files(review)[0]
        self.assertEqual(file["readiness"], "PDF_NOT_READY")
        self.assertEqual(file["pages_read"], 1)
        self.assertEqual(file["page_read_errors"], 1)
        self.assertEqual(rows[0]["error_stage"], "page_read")
        self.assertFalse(rows[0]["page_read_ok"])
        self.assertEqual(rows[1]["class"], "digital")

    def test_classification_failure_preserves_text_readiness(self):
        (self.inputs / "one.pdf").write_bytes(sample_pdf())
        with mock.patch.object(app, "classify_page", side_effect=ValueError("drawing failure")):
            rows, review = self.scan()
        self.assertEqual(rows[0]["error_stage"], "classification")
        self.assertTrue(rows[0]["has_extractable_text"])
        self.assertGreater(rows[0]["text_chars"], 40)
        file = self.files(review)[0]
        self.assertEqual(file["readiness"], "READY_WITH_TEXT")
        self.assertEqual(file["classification_errors"], 1)
        self.assertEqual(review.stats["errors"], 1)

    def test_empty_password_and_required_password_are_distinct(self):
        with fitz.open(stream=sample_pdf(), filetype="pdf") as doc:
            for name, password in (("empty.pdf", ""), ("locked.pdf", "secret")):
                (self.inputs / name).write_bytes(
                    doc.tobytes(
                        encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw="owner"
                    )
                )
        rows, review = self.scan()
        files = {r["file_name"]: r for r in self.files(review)}
        self.assertTrue(files["empty.pdf"]["encrypted"])
        self.assertFalse(files["empty.pdf"]["password_required"])
        self.assertEqual(files["empty.pdf"]["readiness"], "READY_WITH_TEXT")
        self.assertTrue(files["locked.pdf"]["password_required"])
        self.assertIn("requires a password", files["locked.pdf"]["file_error"])
        self.assertEqual(files["locked.pdf"]["readiness"], "PDF_NOT_READY")

    def test_pdf_limits_keep_one_overview_row(self):
        (self.inputs / "large.pdf").write_bytes(sample_pdf(2))
        for settings in ({"max_pages": 1}, {"max_pdf_gib": 1 / 1024**3}):
            rows, review = self.scan(**settings)
            self.assertEqual(len(rows), 1)
            self.assertEqual(len(self.files(review)), 1)
            self.assertEqual(self.files(review)[0]["readiness"], "PDF_NOT_READY")
            self.assertEqual(review.stats["errors"], 1)

    def test_nonfatal_parser_warnings_are_isolated(self):
        (self.inputs / "a.pdf").write_bytes(sample_pdf())
        (self.inputs / "b.pdf").write_bytes(sample_pdf())
        old_errors = fitz.TOOLS.mupdf_display_errors()
        with mock.patch.object(fitz.TOOLS, "mupdf_warnings", side_effect=["parser warning", ""]):
            rows, review = self.scan()
        files = self.files(review)
        self.assertEqual([r["parser_warning_count"] for r in files], [1, 0])
        self.assertEqual(review.stats["errors"], 0)
        self.assertEqual(review.stats["parser_warnings"], 1)
        self.assertTrue(all(r["readable"] for r in files))
        self.assertEqual(fitz.TOOLS.mupdf_display_errors(), old_errors)

    def test_missing_eof_is_information_not_a_readiness_failure(self):
        (self.inputs / "tail.pdf").write_bytes(sample_pdf().replace(b"%%EOF", b"     "))
        rows, review = self.scan()
        file = self.files(review)[0]
        self.assertTrue(file["header_present"])
        self.assertFalse(file["eof_present"])
        self.assertTrue(file["readable"])
        self.assertEqual(review.stats["errors"], 0)

    def test_real_parser_repair_and_clean_following_pdf(self):
        data = sample_pdf()
        data = data[: data.rfind(b"startxref")] + b"startxref\n0\n%%EOF\n"
        (self.inputs / "a-repaired.pdf").write_bytes(data)
        (self.inputs / "b-clean.pdf").write_bytes(sample_pdf())
        rows, review = self.scan()
        files = self.files(review)
        self.assertTrue(files[0]["parser_repaired"])
        self.assertGreater(files[0]["parser_warning_count"], 0)
        self.assertFalse(files[1]["parser_repaired"])
        self.assertEqual(files[1]["parser_warning_count"], 0)
        self.assertTrue(all(r["readable"] for r in files))
        self.assertEqual(review.stats["errors"], 0)

    def test_real_selected_pdf_crc_failure_does_not_skip_other_members(self):
        payload = sample_pdf()
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_STORED) as archive:
            archive.writestr("broken.pdf", payload)
            archive.writestr("good.pdf", payload)
        data = bytearray(stream.getvalue())
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            info = archive.infolist()[0]
            offset = info.header_offset + 30 + len(info.filename.encode("utf-8"))
        data[offset] ^= 1
        (self.inputs / "one.zip").write_bytes(data)
        rows, review = self.scan(check_zip_crc=True)
        self.assertEqual(review.stats["pdfs_found"], 2)
        self.assertEqual(len(self.files(review)), 2)
        self.assertEqual(review.stats["errors"], 1)
        self.assertEqual(
            [r["readiness"] for r in self.files(review)], ["PDF_NOT_READY", "READY_WITH_TEXT"]
        )
        self.assertEqual(review.stats["zips_full_crc_failed"], 1)
        self.assertEqual(review.stats["zip_members_crc_checked"], 1)

    def test_full_zip_crc_checks_skipped_files_without_creating_pdf_errors(self):
        payload = b"ordinary skipped member"
        data = legacy_zip("folder/note.txt", payload, "cp437")
        offset = 30 + len("folder/note.txt")
        data = data[:offset] + bytes([data[offset] ^ 1]) + data[offset + 1 :]
        (self.inputs / "bad-crc.zip").write_bytes(data)
        rows, review = self.scan()
        self.assertEqual(rows, [])
        self.assertEqual(review.stats["errors"], 0)
        rows, review = self.scan(check_zip_crc=True)
        self.assertEqual(review.stats["errors"], 1)
        self.assertEqual(review.stats["pdfs_found"], 0)
        self.assertEqual(self.files(review), [])
        self.assertEqual(rows[0]["object_type"], "file")
        self.assertEqual(rows[0]["levels"], ["bad-crc.zip", "folder", "note.txt"])
        self.assertEqual(review.stats["zips_full_crc_failed"], 1)

    def test_full_zip_crc_reconciles_non_pdf_archives(self):
        (self.inputs / "one.zip").write_bytes(zip_bytes([("leaf.pdf", sample_pdf())]))
        (self.inputs / "two.zip").write_bytes(zip_bytes([("note.txt", b"skip")]))
        (self.inputs / "three.zip").write_bytes(zip_bytes([]))
        rows, review = self.scan(check_zip_crc=True)
        self.assertEqual(review.stats["zips_opened"], 3)
        self.assertEqual(review.stats["zips_with_direct_pdf"], 1)
        self.assertEqual(review.stats["zips_without_direct_pdf"], 2)
        self.assertEqual(review.stats["zips_full_crc_passed"], 3)
        self.assertEqual(review.stats["zip_members_crc_checked"], 2)
        self.assertEqual(review.stats["skipped_files"], 1)
        self.assertEqual(review.stats["errors"], 0)

    def test_pdf_member_read_failure_is_in_file_overview(self):
        (self.inputs / "one.zip").write_bytes(zip_bytes([("leaf.pdf", sample_pdf())]))
        with mock.patch.object(app.Review, "copy_member", side_effect=ValueError("CRC error")):
            rows, review = self.scan(check_zip_crc=True)
        self.assertEqual(len(self.files(review)), 1)
        self.assertEqual(review.stats["pdfs_found"], 1)
        self.assertEqual(self.files(review)[0]["readiness"], "PDF_NOT_READY")
        self.assertEqual(review.stats["zips_full_crc_failed"], 1)

    def test_nonfinite_limits_are_rejected(self):
        for option in ("--max-pdf-gib", "--max-extracted-gib"):
            for value in ("nan", "inf", "0", "-1"):
                with (
                    mock.patch.object(app, "PROJECT_DIR", self.root),
                    self.assertRaises(ValueError),
                ):
                    app.run(["--root", str(self.inputs), option, value])

    def test_timestamp_output_and_both_sheets_are_verified(self):
        (self.inputs / "one.pdf").write_bytes(sample_pdf(2))
        output = self.root / "review.xlsx"
        with mock.patch.object(app, "PROJECT_DIR", self.root), redirect_stdout(io.StringIO()):
            for _ in range(2):
                self.assertEqual(
                    app.run(
                        ["--root", str(self.inputs), "--output", str(output), "--timestamp-output"]
                    ),
                    0,
                )
        outputs = list(self.root.glob("review-*.xlsx"))
        self.assertEqual(len(outputs), 2)
        for output in outputs:
            workbook = load_workbook(output)
            try:
                self.assertEqual(workbook.sheetnames, ["pdf_pages", "pdf_files"])
                file = workbook["pdf_files"]
                self.assertEqual(file["B2"].value, "Table: PDF file readiness")
                self.assertEqual(file.max_row, 7)
                self.assertEqual(file["C6"].fill.fgColor.rgb, "FF00B0F0")
                summary = json.loads(output.with_name(output.stem + "-summary.json").read_text())
                self.assertEqual(summary["pdf_file_rows"], 1)
                self.assertEqual(summary["pages_with_text"], 2)
                self.assertEqual(summary["readiness"], {"READY_WITH_TEXT": 1})
            finally:
                workbook.close()

    def test_loose_case_insensitive_and_skips(self):
        (self.inputs / "document.PdF").write_bytes(sample_pdf(2))
        (self.inputs / "note.txt").write_text("ordinary text")
        (self.inputs / "empty.zip").write_bytes(zip_bytes([]))
        (self.inputs / "others.zip").write_bytes(zip_bytes([("note.doc", b"other file")]))
        rows, review = self.scan()
        self.assertEqual(len(rows), 2)
        self.assertEqual([r["page_number"] for r in rows], [1, 2])
        self.assertEqual(review.stats["errors"], 0)
        self.assertEqual(review.stats["skipped_files"], 2)

    def test_mixed_nested_archives(self):
        inner = zip_bytes([("Bottom/leaf.PDF", sample_pdf(2)), ("ignored.txt", b"skip")])
        rar = rar_bytes([("Middle/inner.zip", inner)])
        seven = self.seven_zip([("Layer/middle.rar", rar)])
        (self.inputs / "outer.zip").write_bytes(zip_bytes([("Top/middle.7z", seven)]))
        rows, review = self.scan()
        self.assertEqual(len(rows), 2)
        self.assertEqual(review.stats["errors"], 0, review.errors.getvalue())
        self.assertEqual(rows[0]["archive_depth"], 4)
        self.assertEqual(
            rows[0]["levels"],
            [
                "outer.zip",
                "Top",
                "middle.7z",
                "Layer",
                "middle.rar",
                "Middle",
                "inner.zip",
                "Bottom",
                "leaf.PDF",
            ],
        )
        self.assertTrue(
            rows[0]["full_path"].endswith(
                "outer.zip::Top/middle.7z::Layer/middle.rar::Middle/inner.zip::Bottom/leaf.PDF"
            )
        )

    def test_reverse_order_and_solid_7z(self):
        inner = self.seven_zip(
            [("one.pdf", sample_pdf()), ("two.pdf", sample_pdf(2)), ("other.txt", b"skip")]
        )
        middle = zip_bytes([("inner.7z", inner)])
        (self.inputs / "outer.7z").write_bytes(
            self.seven_zip([("nested.zip", middle), ("direct.pdf", sample_pdf())])
        )
        rows, review = self.scan()
        self.assertEqual(len(rows), 4)
        self.assertEqual(review.stats["errors"], 0, review.errors.getvalue())
        self.assertEqual(review.stats["pdfs_found"], 3)

    def test_long_unicode_and_literal_names(self):
        name = "Каталог/" + "Документ東京" * 60 + ".PDF"
        (self.inputs / "unicode.zip").write_bytes(
            zip_bytes([(name, sample_pdf()), ("=SUM(1,2).pdf", sample_pdf())])
        )
        rows, review = self.scan()
        self.assertEqual(review.stats["errors"], 0)
        self.assertEqual({r["file_name"] for r in rows}, {name.split("/")[-1], "=SUM(1,2).pdf"})
        spool = self.root / "rows.jsonl"
        spool.write_text(
            "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8"
        )
        output = self.root / "review.xlsx"
        pdf_spool = self.root / "pdfs.jsonl"
        pdf_spool.write_text(review.pdfs.getvalue(), encoding="utf-8")
        app.build_xlsx(
            spool,
            output,
            len(rows),
            review.stats["max_level"],
            pdf_spool,
            review.stats["pdf_file_rows"],
        )
        workbook = load_workbook(output)
        try:
            sheet = workbook["pdf_pages"]
            self.assertEqual(sheet["B6"].value, "№пп")
            self.assertEqual(sheet["B7"].value, "=ROW()-ROW($B$6)")
            self.assertEqual(sheet["C6"].fill.fgColor.rgb, "FF00B0F0")
            self.assertEqual(sheet.sheet_view.zoomScale, 85)
            self.assertIs(sheet.sheet_view.showGridLines, False)
            self.assertIsNone(sheet.freeze_panes)
            self.assertTrue(sheet.auto_filter.ref.startswith("B6:"))
            self.assertEqual(sheet.cell(6, sheet.max_column).value, "last")
            values = {
                c.value for row in sheet.iter_rows(min_row=7) for c in row if c.data_type == "s"
            }
            self.assertIn(name.split("/")[-1], values)
            self.assertIn("=SUM(1,2).pdf", values)
            self.assertFalse(
                any(c.data_type == "f" for row in sheet.iter_rows(min_row=7) for c in row[2:])
            )
            file_sheet = workbook["pdf_files"]
            self.assertIn(
                name.split("/")[-1],
                {
                    c.value
                    for row in file_sheet.iter_rows(min_row=7)
                    for c in row
                    if c.data_type == "s"
                },
            )
            self.assertFalse(
                any(c.data_type == "f" for row in file_sheet.iter_rows(min_row=7) for c in row[2:])
            )
        finally:
            workbook.close()

    def test_legacy_zip_encodings_and_unicode_metadata(self):
        for encoding in ("cp866", "cp1251"):
            with zipfile.ZipFile(
                io.BytesIO(legacy_zip("Отчёт.pdf", sample_pdf(), encoding))
            ) as archive:
                self.assertEqual(app.zip_name(archive.infolist()[0], encoding), "Отчёт.pdf")
        data = legacy_zip("plain.pdf", sample_pdf(), "cp437", unicode_name="Документ東京.pdf")
        (self.inputs / "extra.zip").write_bytes(data)
        rows, review = self.scan()
        self.assertEqual(review.stats["errors"], 0)
        self.assertEqual(rows[0]["file_name"], "Документ東京.pdf")

    def test_duplicate_zip_names_remain_separate_pdfs(self):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            data = zip_bytes([("same.pdf", sample_pdf()), ("same.pdf", sample_pdf(2))])
        (self.inputs / "duplicates.zip").write_bytes(data)
        rows, review = self.scan()
        self.assertEqual(review.stats["errors"], 0)
        self.assertEqual(len(rows), 3)
        self.assertEqual(len({r["pdf_id"] for r in rows}), 2)

    def test_errors_are_visible_and_password_pdf(self):
        (self.inputs / "broken.zip").write_bytes(b"broken")
        (self.inputs / "broken.pdf").write_bytes(b"not PDF")
        with fitz.open(stream=sample_pdf(), filetype="pdf") as doc:
            locked = doc.tobytes(
                encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="sample", owner_pw="owner"
            )
        (self.inputs / "locked.pdf").write_bytes(locked)
        rows, review = self.scan()
        self.assertEqual(review.stats["errors"], 3)
        self.assertTrue(all(r["class"] == "error" and r["page_number"] is None for r in rows))
        self.assertEqual({r["object_type"] for r in rows}, {"pdf", "archive"})

    def test_page_failure_does_not_skip_later_pages(self):
        (self.inputs / "document.pdf").write_bytes(sample_pdf(2))
        original = app.classify_page

        def classify(page, **kwargs):
            if page.number == 0:
                raise ValueError("synthetic page failure")
            return original(page, **kwargs)

        with mock.patch.object(app, "classify_page", side_effect=classify):
            rows, review = self.scan()
        self.assertEqual([r["class"] for r in rows], ["error", "digital"])
        self.assertEqual([r["page_number"] for r in rows], [1, 2])

    def test_depth_and_decompression_limits(self):
        inner = zip_bytes([("leaf.pdf", sample_pdf())])
        (self.inputs / "outer.zip").write_bytes(zip_bytes([("inner.zip", inner)]))
        rows, review = self.scan(max_depth=1)
        self.assertEqual(review.stats["errors"], 1)
        self.assertEqual(rows[0]["object_type"], "archive")
        rows, review = self.scan(max_extracted_gib=1 / 1024**3)
        self.assertEqual(review.stats["errors"], 1)
        self.assertEqual(rows[0]["class"], "error")

    def test_member_names_never_direct_extraction(self):
        (self.inputs / "paths.zip").write_bytes(
            zip_bytes([("../outside.pdf", sample_pdf()), ("/absolute.pdf", sample_pdf())])
        )
        rows, review = self.scan()
        self.assertEqual(len(rows), 2)
        self.assertEqual(review.stats["errors"], 0)
        self.assertFalse((self.root / "outside.pdf").exists())

    def test_windows_path_conversion(self):
        with (
            mock.patch.object(app.os, "name", "nt"),
            mock.patch.object(app.os.path, "abspath", side_effect=lambda p: str(p)),
        ):
            self.assertEqual(app.fs_path(r"C:\Data\file.pdf"), r"\\?\C:\Data\file.pdf")
            self.assertEqual(
                app.fs_path(r"\\server\share\file.pdf"), r"\\?\UNC\server\share\file.pdf"
            )
            self.assertEqual(app.fs_path(r"\\?\C:\Data\file.pdf"), r"\\?\C:\Data\file.pdf")

    def test_run_writes_summary_and_returns_error_status(self):
        (self.inputs / "document.pdf").write_bytes(sample_pdf(2))
        runtime_project = self.root / "project"
        runtime_project.mkdir()
        output = self.root / "out" / "review.xlsx"
        with mock.patch.object(app, "PROJECT_DIR", runtime_project), redirect_stdout(io.StringIO()):
            self.assertEqual(app.run(["--root", str(self.inputs), "--output", str(output)]), 0)
            (self.inputs / "bad.pdf").write_bytes(b"bad")
            self.assertEqual(app.run(["--root", str(self.inputs), "--output", str(output)]), 2)
        summary = json.loads(output.with_name("review-summary.json").read_text())
        self.assertEqual(summary["pages"], 2)
        self.assertEqual(summary["rows"], 3)
        self.assertEqual(summary["errors"], 1)

    def test_physical_paths_over_260_characters(self):
        directory = self.inputs
        for i in range(4):
            directory = directory / (f"level-{i}-" + "a" * 70)
            os.makedirs(app.fs_path(directory), exist_ok=True)
        with open(app.fs_path(directory / "document.PDF"), "wb") as handle:
            handle.write(sample_pdf())
        rows, review = self.scan()
        self.assertEqual(len(rows), 1)
        self.assertEqual(review.stats["errors"], 0)

    def test_ten_archive_levels(self):
        data = zip_bytes([("bottom.pdf", sample_pdf())])
        for index in range(9):
            data = zip_bytes([(f"folder-{index}/inner.zip", data)])
        (self.inputs / "outer.zip").write_bytes(data)
        rows, review = self.scan()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["archive_depth"], 10)
        self.assertEqual(review.stats["errors"], 0)

    def test_empty_input_is_successful(self):
        output = self.root / "empty.xlsx"
        with mock.patch.object(app, "PROJECT_DIR", self.root), redirect_stdout(io.StringIO()):
            self.assertEqual(app.run(["--root", str(self.inputs), "--output", str(output)]), 0)
        workbook = load_workbook(output, read_only=True)
        try:
            self.assertEqual(sum(1 for _ in workbook["pdf_pages"].iter_rows()), 6)
        finally:
            workbook.close()


class ClassifierTests(unittest.TestCase):
    def classify(self, page, **kwargs):
        settings = {
            "text_threshold": 40,
            "img_cover_threshold": 0.9,
            "min_img_frac": 0.005,
            "vector_threshold": 100,
        }
        settings.update(kwargs)
        return app.classify_page(page, **settings)

    def image(self):
        pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 4, 4), False)
        pixmap.clear_with(128)
        return pixmap.tobytes("png")

    def test_all_classes_and_vector_boundary(self):
        with fitz.open() as doc:
            page = doc.new_page(width=300, height=600)
            self.assertEqual(self.classify(page)[0], "other")
            for i in range(101):
                page.draw_rect(fitz.Rect(i, 10, i + 1, 11))
            self.assertEqual(self.classify(page)[0], "vector")
            self.assertEqual(self.classify(page, vector_threshold=101)[0], "other")
            page.insert_text(
                (20, 50), "This page contains more than forty extracted text characters."
            )
            self.assertEqual(self.classify(page)[0], "digital")
            page.insert_image(page.rect, stream=self.image(), keep_proportion=False)
            self.assertEqual(self.classify(page)[0], "ocr")
            page = doc.new_page(width=300, height=600)
            page.insert_image(page.rect, stream=self.image(), keep_proportion=False)
            self.assertEqual(self.classify(page)[0], "scan")

    def test_rotated_and_cropped_full_page_scan(self):
        with fitz.open() as doc:
            page = doc.new_page(width=300, height=600)
            page.insert_image(page.rect, stream=self.image(), keep_proportion=False)
            for rotation in (0, 90, 180, 270):
                page.set_rotation(rotation)
                self.assertEqual(self.classify(page)[:3], ("scan", 0, 1.0))
            page.set_rotation(0)
            page.set_cropbox(fitz.Rect(20, 30, 280, 550))
            self.assertEqual(self.classify(page)[:3], ("scan", 0, 1.0))

    def test_overlaps_are_counted_once(self):
        self.assertEqual(app._union_area([(0, 0, 10, 10), (5, 0, 15, 10)]), 150)
        with fitz.open() as doc:
            page = doc.new_page(width=300, height=300)
            for _ in range(2):
                page.insert_image(
                    fitz.Rect(0, 0, 150, 300), stream=self.image(), keep_proportion=False
                )
            self.assertEqual(self.classify(page)[2], 0.5)


if __name__ == "__main__":
    unittest.main()
